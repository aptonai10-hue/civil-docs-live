from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path('/home/ubuntu/deliverables/CivilDocs_90_Day_Revenue_and_Expense_Projection_2026-08-20.xlsx')
OUT.parent.mkdir(parents=True, exist_ok=True)

DARK_GREEN = '135B44'
LIGHT_GREEN = 'CFE9E0'
BLUE = '0000FF'
BLACK = '000000'
GREEN = '008000'
WHITE = 'FFFFFF'
MEDIUM_GREEN = Side(style='medium', color=DARK_GREEN)
CURRENCY = '"ZMW "#,##0.0;("ZMW "#,##0.0);-'
NUMBER = '#,##0.0;(#,##0.0);-'
PERCENT = '#,##0.0%'

wb = Workbook()
ass = wb.active
ass.title = 'Assumptions'
model = wb.create_sheet('Revenue and P&L')
expenses = wb.create_sheet('Expense Detail')
checks = wb.create_sheet('Checks')

def set_title(ws, title, subtitle, unit):
    for col in range(3, 7):
        cell = ws.cell(3, col)
        cell.fill = PatternFill('solid', fgColor=DARK_GREEN)
        cell.font = Font(color=WHITE, bold=True, size=16)
    ws['C3'] = title
    ws['C5'] = subtitle
    ws['C5'].font = Font(bold=True, size=11)
    ws['C6'] = unit
    ws['C6'].font = Font(italic=True)
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20
    ws.sheet_view.showGridLines = False

def set_section(ws, row, text, end_col=6):
    for col in range(3, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill('solid', fgColor=LIGHT_GREEN)
        cell.font = Font(bold=True, color=BLACK)
    ws.cell(row, 3).value = text

def label(ws, cell, text, italic=False, bold=False):
    ws[cell] = text
    ws[cell].font = Font(italic=italic, bold=bold, color=BLACK)

def input_cell(ws, cell, value, fmt=NUMBER, comment_text='Source: CivilDocs management base-case assumption, 2026-08-20. Replace with observed or supplier data before paid launch.'):
    target = ws[cell]
    target.value = value
    target.number_format = fmt
    target.font = Font(color=BLUE)
    target.alignment = Alignment(horizontal='right')
    target.comment = Comment(comment_text, 'Manus AI')

def formula_cell(ws, cell, formula, fmt=NUMBER, cross_sheet=False):
    target = ws[cell]
    target.value = formula
    target.number_format = fmt
    target.font = Font(color=GREEN if cross_sheet else BLACK)
    target.alignment = Alignment(horizontal='right')

def subtotal(ws, row, start_col=3, end_col=6, final=False):
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color=cell.font.color.rgb if cell.font.color and cell.font.color.type == 'rgb' else BLACK)
        cell.border = Border(top=MEDIUM_GREEN, bottom=Side(style='double', color=DARK_GREEN) if final else Side(style=None))

def autosize(ws, max_col):
    for col in range(3, max_col + 1):
        width = 12
        for row in range(1, ws.max_row + 1):
            value = ws.cell(row, col).value
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 54))
        ws.column_dimensions[get_column_letter(col)].width = width

def configure(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.column >= 4 and isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal='right')
    autosize(ws, max(6, ws.max_column))
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'B2:{get_column_letter(ws.max_column)}{ws.max_row}'
    ws.oddFooter.center.text = f'{ws.title} — Page &P of &N'

set_title(ass, 'CivilDocs 90-Day Financial Model — Assumptions', 'Base case for subscription billings, operating expense and cash profit/loss', '(ZMW unless stated; blue cells are editable management assumptions)')
set_section(ass, 8, 'Pricing and Plan Structure')
label(ass, 'C9', 'Founding Engineer — monthly price')
input_cell(ass, 'D9', 199, CURRENCY, 'Source: CivilDocs commercial pricing recommendation, 2026-08-20. Introductory monthly price for first three paid months.')
label(ass, 'C10', 'Professional — monthly price')
input_cell(ass, 'D10', 299, CURRENCY, 'Source: CivilDocs commercial pricing recommendation, 2026-08-20. Standard proposed per-user price.')
label(ass, 'C11', 'Firm — monthly price')
input_cell(ass, 'D11', 1199, CURRENCY, 'Source: CivilDocs commercial pricing recommendation, 2026-08-20. Five-user Firm bundle price.')
label(ass, 'C12', 'Firm included user seats')
input_cell(ass, 'D12', 5, NUMBER, 'Source: CivilDocs commercial pricing recommendation, 2026-08-20. Included named user seats.')
set_section(ass, 14, 'Pilot Conversion and Acquisition')
label(ass, 'C15', 'Invited pilot engineers')
input_cell(ass, 'D15', 6, NUMBER, 'Source: User-provided project context. Six experienced civil engineers are the proposed pilot cohort.')
label(ass, 'C16', 'Pilot-to-Founding conversion rate')
input_cell(ass, 'D16', 4 / 6, PERCENT, 'Source: Management base-case assumption, 2026-08-20. Four of six invited pilots convert in Month 1. Not observed performance.')
label(ass, 'C17', 'Founding Engineer seats at Month 1 start', bold=True)
formula_cell(ass, 'D17', '=ROUND(D15*D16,0)')
label(ass, 'C18', 'Net new Founding Engineer seats — Month 2')
input_cell(ass, 'D18', 1)
label(ass, 'C19', 'Net new Founding Engineer seats — Month 3')
input_cell(ass, 'D19', 1)
label(ass, 'C20', 'New Professional seats — Month 2')
input_cell(ass, 'D20', 2)
label(ass, 'C21', 'New Professional seats — Month 3')
input_cell(ass, 'D21', 3)
label(ass, 'C22', 'New Firm accounts — Month 3')
input_cell(ass, 'D22', 1)
set_section(ass, 24, 'Fair-Use and Variable-Cost Drivers')
label(ass, 'C25', 'Included AI document generations per paid user/month')
input_cell(ass, 'D25', 30, NUMBER, 'Source: CivilDocs fair-use policy draft, 2026-08-20. Paid launch requires account-level metering before this allowance is enforceable.')
label(ass, 'C26', 'Included Firm AI generations per month')
formula_cell(ass, 'D26', '=D12*D25')
label(ass, 'C27', 'Assumed included-capacity utilisation')
input_cell(ass, 'D27', 0.25, PERCENT, 'Source: Management base-case assumption, 2026-08-20. 25% of included generation capacity is used. Replace with observed usage.')
label(ass, 'C28', 'AI model cost per successful generation')
input_cell(ass, 'D28', 2.5, CURRENCY, 'Source: Management base-case placeholder, 2026-08-20. Replace with actual Gemini/model-provider blended cost including retry waste.')
label(ass, 'C29', 'Payment-processing fee as % of collections')
input_cell(ass, 'D29', 0.035, PERCENT, 'Source: Management base-case placeholder, 2026-08-20. Replace with the contracted payment-provider fee and local tax treatment.')
set_section(ass, 31, 'Fixed Operating Expense Drivers')
label(ass, 'C32', 'Hosting and infrastructure — monthly')
input_cell(ass, 'D32', 250, CURRENCY, 'Source: Management base-case placeholder, 2026-08-20. Replace with production hosting, monitoring and storage invoices.')
label(ass, 'C33', 'Support and administration labour — monthly')
input_cell(ass, 'D33', 1000, CURRENCY, 'Source: Management base-case placeholder, 2026-08-20. Represents part-time support and administration only; founder compensation is excluded.')
label(ass, 'C34', 'Sales and customer acquisition — monthly')
input_cell(ass, 'D34', 600, CURRENCY, 'Source: Management base-case placeholder, 2026-08-20. Includes launch outreach and acquisition activity, excluding commissions.')
label(ass, 'C35', 'Software and communications — monthly')
input_cell(ass, 'D35', 150, CURRENCY, 'Source: Management base-case placeholder, 2026-08-20. Includes communications and other SaaS tools.')
label(ass, 'C36', 'Legal and compliance setup — Month 1 only')
input_cell(ass, 'D36', 1500, CURRENCY, 'Source: Management base-case placeholder, 2026-08-20. Budget for lawyer review, privacy and terms preparation; confirm with actual provider quotes.')
set_section(ass, 38, 'Model Boundaries')
for cell, text in {
    'D39': 'VAT, income tax, founder/executive compensation, capital expenditure, debt service, refunds, bad debt and working-capital timing are excluded.',
    'D40': 'The model assumes monthly billing is collected in the same month; it is a cash operating-planning view, not audited financial statements.',
    'D41': 'All blue expense assumptions are placeholders until supplier invoices, payment contracts, model usage and staffing commitments are known.'
}.items():
    ass[cell] = text
    ass[cell].font = Font(color=BLUE)
    ass[cell].alignment = Alignment(wrap_text=True, vertical='top')
    ass[cell].comment = Comment('Source: CivilDocs model scope definition, 2026-08-20.', 'Manus AI')
subtotal(ass, 17, 3, 4)
subtotal(ass, 26, 3, 4)

set_title(expenses, 'CivilDocs 90-Day Expense Detail', 'Fixed operating-cost schedule linked to editable assumptions', '(ZMW; excluding variable AI-model and payment-processing costs)')
for col, heading in zip(range(4, 7), ['Month 1', 'Month 2', 'Month 3']):
    expenses.cell(8, col).value = heading
    expenses.cell(8, col).font = Font(bold=True)
    expenses.cell(8, col).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
    expenses.cell(8, col).alignment = Alignment(horizontal='center')
set_section(expenses, 10, 'Fixed Operating Expenses')
fixed_rows = [
    (11, 'Hosting and infrastructure', '=Assumptions!$D$32'),
    (12, 'Support and administration labour', '=Assumptions!$D$33'),
    (13, 'Sales and customer acquisition', '=Assumptions!$D$34'),
    (14, 'Software and communications', '=Assumptions!$D$35'),
]
for row, name, formula in fixed_rows:
    label(expenses, f'C{row}', name)
    for col in 'DEF':
        formula_cell(expenses, f'{col}{row}', formula, CURRENCY, cross_sheet=True)
label(expenses, 'C15', 'Legal and compliance setup')
formula_cell(expenses, 'D15', '=Assumptions!$D$36', CURRENCY, cross_sheet=True)
formula_cell(expenses, 'E15', '=0', CURRENCY)
formula_cell(expenses, 'F15', '=0', CURRENCY)
label(expenses, 'C16', 'Total fixed operating expenses', bold=True)
for col in 'DEF':
    formula_cell(expenses, f'{col}16', f'=SUM({col}11:{col}15)', CURRENCY)
subtotal(expenses, 16, 3, 6, final=True)
label(expenses, 'C18', 'Expense boundary note')
expenses['D18'] = 'Founder compensation, taxes, refunds, debt service, capital expenditure and working-capital timing are excluded. Add them before using this as a full cash-flow forecast.'
expenses['D18'].font = Font(color=BLUE)
expenses['D18'].alignment = Alignment(wrap_text=True, vertical='top')
expenses['D18'].comment = Comment('Source: CivilDocs model scope definition, 2026-08-20.', 'Manus AI')

set_title(model, 'CivilDocs 90-Day Revenue, Expense and Profit/Loss', 'Base-case monthly gross billings, cash operating expense and operating profit/loss', '(ZMW unless stated; negative values represent cash operating losses)')
for col, heading in zip(range(4, 7), ['Month 1', 'Month 2', 'Month 3']):
    model.cell(8, col).value = heading
    model.cell(8, col).font = Font(bold=True)
    model.cell(8, col).fill = PatternFill('solid', fgColor=LIGHT_GREEN)
    model.cell(8, col).alignment = Alignment(horizontal='center')
set_section(model, 10, 'Paid Subscription Base')
label(model, 'C11', 'Founding Engineer paid seats')
formula_cell(model, 'D11', '=Assumptions!$D$17', cross_sheet=True)
formula_cell(model, 'E11', '=D11+Assumptions!$D$18', cross_sheet=True)
formula_cell(model, 'F11', '=E11+Assumptions!$D$19', cross_sheet=True)
label(model, 'C12', 'Professional paid seats')
formula_cell(model, 'D12', '=0')
formula_cell(model, 'E12', '=Assumptions!$D$20', cross_sheet=True)
formula_cell(model, 'F12', '=E12+Assumptions!$D$21', cross_sheet=True)
label(model, 'C13', 'Firm accounts')
formula_cell(model, 'D13', '=0')
formula_cell(model, 'E13', '=0')
formula_cell(model, 'F13', '=Assumptions!$D$22', cross_sheet=True)
set_section(model, 15, 'Gross Subscription Billings')
label(model, 'C16', 'Founding Engineer revenue')
for col in 'DEF': formula_cell(model, f'{col}16', f'={col}11*Assumptions!$D$9', CURRENCY, cross_sheet=True)
label(model, 'C17', 'Professional revenue')
for col in 'DEF': formula_cell(model, f'{col}17', f'={col}12*Assumptions!$D$10', CURRENCY, cross_sheet=True)
label(model, 'C18', 'Firm revenue')
for col in 'DEF': formula_cell(model, f'{col}18', f'={col}13*Assumptions!$D$11', CURRENCY, cross_sheet=True)
label(model, 'C19', 'Total gross subscription billings', bold=True)
for col in 'DEF': formula_cell(model, f'{col}19', f'=SUM({col}16:{col}18)', CURRENCY)
subtotal(model, 19, 3, 6, final=True)
set_section(model, 21, 'Usage and Variable Expenses')
label(model, 'C22', 'Included AI generation capacity')
for col in 'DEF': formula_cell(model, f'{col}22', f'=({col}11+{col}12+( {col}13*Assumptions!$D$12))*Assumptions!$D$25', cross_sheet=True)
label(model, 'C23', 'Assumed successful AI generations')
for col in 'DEF': formula_cell(model, f'{col}23', f'={col}22*Assumptions!$D$27', cross_sheet=True)
label(model, 'C24', 'AI model provider expense')
for col in 'DEF': formula_cell(model, f'{col}24', f'={col}23*Assumptions!$D$28', CURRENCY, cross_sheet=True)
label(model, 'C25', 'Payment-processing expense')
for col in 'DEF': formula_cell(model, f'{col}25', f'={col}19*Assumptions!$D$29', CURRENCY, cross_sheet=True)
label(model, 'C26', 'Total variable expenses', bold=True)
for col in 'DEF': formula_cell(model, f'{col}26', f'=SUM({col}24:{col}25)', CURRENCY)
subtotal(model, 26, 3, 6)
set_section(model, 28, 'Cash Operating Profit / Loss')
label(model, 'C29', 'Contribution after variable expenses', bold=True)
for col in 'DEF': formula_cell(model, f'{col}29', f'={col}19-{col}26', CURRENCY)
label(model, 'C30', 'Fixed operating expenses')
for col in 'DEF': formula_cell(model, f'{col}30', f'=\'Expense Detail\'!{col}16', CURRENCY, cross_sheet=True)
label(model, 'C31', 'Total cash operating expenses')
for col in 'DEF': formula_cell(model, f'{col}31', f'={col}26+{col}30', CURRENCY)
label(model, 'C32', 'Cash operating profit / (loss)', bold=True)
for col in 'DEF': formula_cell(model, f'{col}32', f'={col}19-{col}31', CURRENCY)
label(model, 'C33', 'Cash operating margin', italic=True)
for col in 'DEF': formula_cell(model, f'{col}33', f'=IFERROR({col}32/{col}19,0)', PERCENT)
subtotal(model, 32, 3, 6, final=True)
set_section(model, 35, '90-Day Summary')
label(model, 'C36', '90-day gross subscription billings', bold=True)
formula_cell(model, 'D36', '=SUM(D19:F19)', CURRENCY)
label(model, 'C37', '90-day cash operating expenses', bold=True)
formula_cell(model, 'D37', '=SUM(D31:F31)', CURRENCY)
label(model, 'C38', '90-day cash operating profit / (loss)', bold=True)
formula_cell(model, 'D38', '=D36-D37', CURRENCY)
label(model, 'C39', 'Month 3 exit monthly recurring revenue')
formula_cell(model, 'D39', '=F19', CURRENCY)
label(model, 'C40', 'Month 3 paid-seat equivalent')
formula_cell(model, 'D40', '=F11+F12+(F13*Assumptions!$D$12)', cross_sheet=True)
label(model, 'C41', '90-day assumed successful AI generations')
formula_cell(model, 'D41', '=SUM(D23:F23)')
subtotal(model, 36, 3, 4)
subtotal(model, 37, 3, 4)
subtotal(model, 38, 3, 4, final=True)

set_title(checks, 'CivilDocs Financial Model — Checks', 'Formula integrity and result checks', '(Values equal to zero indicate a passing arithmetic check)')
set_section(checks, 8, 'Revenue and Expense Checks', 5)
for row, label_text, formula in [
    (9, 'Month 1 revenue check', '=SUM(\'Revenue and P&L\'!D16:D18)-\'Revenue and P&L\'!D19'),
    (10, 'Month 2 revenue check', '=SUM(\'Revenue and P&L\'!E16:E18)-\'Revenue and P&L\'!E19'),
    (11, 'Month 3 revenue check', '=SUM(\'Revenue and P&L\'!F16:F18)-\'Revenue and P&L\'!F19'),
    (12, 'Month 1 expense check', '=\'Revenue and P&L\'!D26+\'Revenue and P&L\'!D30-\'Revenue and P&L\'!D31'),
    (13, 'Month 2 expense check', '=\'Revenue and P&L\'!E26+\'Revenue and P&L\'!E30-\'Revenue and P&L\'!E31'),
    (14, 'Month 3 expense check', '=\'Revenue and P&L\'!F26+\'Revenue and P&L\'!F30-\'Revenue and P&L\'!F31'),
    (15, '90-day profit/loss check', '=SUM(\'Revenue and P&L\'!D32:F32)-\'Revenue and P&L\'!D38'),
]:
    label(checks, f'C{row}', label_text)
    formula_cell(checks, f'D{row}', formula, CURRENCY, cross_sheet=True)
set_section(checks, 17, 'Base-Case Outputs', 5)
for row, label_text, formula, fmt in [
    (18, '90-day gross subscription billings', '=\'Revenue and P&L\'!D36', CURRENCY),
    (19, '90-day cash operating expenses', '=\'Revenue and P&L\'!D37', CURRENCY),
    (20, '90-day cash operating profit / (loss)', '=\'Revenue and P&L\'!D38', CURRENCY),
    (21, 'Month 3 exit monthly recurring revenue', '=\'Revenue and P&L\'!D39', CURRENCY),
    (22, 'Month 3 paid-seat equivalent', '=\'Revenue and P&L\'!D40', NUMBER),
]:
    label(checks, f'C{row}', label_text)
    formula_cell(checks, f'D{row}', formula, fmt, cross_sheet=True)
label(checks, 'C24', 'Interpretation')
checks['D24'] = 'This is a planning model. Replace blue assumptions with actual contracts, invoices and observed usage. It is not a tax, accounting or audited profitability statement.'
checks['D24'].font = Font(color=BLUE)
checks['D24'].alignment = Alignment(wrap_text=True, vertical='top')
checks['D24'].comment = Comment('Source: CivilDocs model scope definition, 2026-08-20.', 'Manus AI')

for ws in wb.worksheets:
    configure(ws)
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = 'auto'
wb.save(OUT)
print(OUT)

